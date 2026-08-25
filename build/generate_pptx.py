#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Rendu PowerPoint du graphe de scène, en formes natives.

    python3 build/generate_pptx.py [--out out/carto.pptx]

Chaque bloc est une forme PowerPoint réelle : sélectionnable, déplaçable,
éditable. Aucune image n'est collée. Les cartes de service portent un lien
hypertexte natif vers la diapositive du service, les graphiques sont des
objets graphiques dont le classeur reste éditable.

Comme generate_html.py, ce script ne connaît aucun contenu métier : il rend
les primitives produites par layout.py.
"""

import argparse
import copy
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Emu, Pt

import layout
import model

RACINE = Path(__file__).resolve().parent.parent
SORTIE = RACINE / "out" / "carto.pptx"
EMU_PAR_POINT = 12700

ALIGNEMENT = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
              "right": PP_ALIGN.RIGHT}
ANCRAGE = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE,
           "bottom": MSO_ANCHOR.BOTTOM}


def emu(points):
    return Emu(int(round(points * EMU_PAR_POINT)))


def rgb(couleur):
    return RGBColor.from_string(couleur.lstrip("#").upper())


def sous_element(parent, balise):
    """Crée l'enfant s'il n'existe pas, et le renvoie."""
    trouve = parent.find(qn(balise))
    if trouve is None:
        trouve = parent.makeelement(qn(balise), {})
        parent.append(trouve)
    return trouve


def sans_ombre(forme):
    """Neutralise l'ombre que le thème applique par défaut.

    `effectLst` doit rester le dernier enfant de `spPr` : remplissage et
    contour sont donc toujours renseignés avant cet appel.
    """
    spPr = forme._element.spPr
    for ancien in spPr.findall(qn("a:effectLst")):
        spPr.remove(ancien)
    spPr.append(spPr.makeelement(qn("a:effectLst"), {}))


# ═════════════════════════════════════════════════════════════════════════════
# Texte
# ═════════════════════════════════════════════════════════════════════════════

def poser_texte(cadre, forme, police):
    """Remplit un cadre de texte selon les attributs d'une primitive."""
    # même règle qu'en HTML : un bloc trop bas pour deux lignes ne se replie
    # jamais, sinon les métriques de Calibri suffisent à le faire déborder
    cadre.word_wrap = not (forme.get("nowrap", False)
                           or forme["h"] < forme["taille"] * 1.9)
    cadre.margin_left = cadre.margin_right = emu(forme.get("pad", 0.0))
    cadre.margin_top = cadre.margin_bottom = 0
    cadre.vertical_anchor = ANCRAGE[forme.get("valign", "middle")]
    paragraphe = cadre.paragraphs[0]
    paragraphe.alignment = ALIGNEMENT[forme["align"]]
    paragraphe.line_spacing = forme.get("interligne", 1.0)
    execution = paragraphe.add_run()
    execution.text = forme["texte"]
    police_run = execution.font
    police_run.name = police
    police_run.size = Pt(forme["taille"])
    police_run.bold = bool(forme.get("gras"))
    police_run.italic = bool(forme.get("italique"))
    police_run.color.rgb = rgb(forme["couleur"])
    # l'autoajustement de PowerPoint réduirait le texte et casserait la
    # géométrie calculée par layout.py
    corps = cadre._txBody.find(qn("a:bodyPr"))
    for balise in ("a:normAutofit", "a:spAutoFit"):
        for noeud in corps.findall(qn(balise)):
            corps.remove(noeud)
    corps.append(corps.makeelement(qn("a:noAutofit"), {}))


# ═════════════════════════════════════════════════════════════════════════════
# Primitives
# ═════════════════════════════════════════════════════════════════════════════

def ajouter_rect(diapo, forme, police, cibles):
    arrondi = forme.get("radius") or 0.0
    genre = MSO_SHAPE.ROUNDED_RECTANGLE if arrondi else MSO_SHAPE.RECTANGLE
    objet = diapo.shapes.add_shape(genre, emu(forme["x"]), emu(forme["y"]),
                                   emu(forme["w"]), emu(forme["h"]))
    if arrondi:
        objet.adjustments[0] = min(0.5, arrondi / min(forme["w"], forme["h"]))

    if forme["fill"]:
        objet.fill.solid()
        objet.fill.fore_color.rgb = rgb(forme["fill"])
    else:
        objet.fill.background()

    if forme["line"] and forme["lw"]:
        objet.line.color.rgb = rgb(forme["line"])
        objet.line.width = emu(forme["lw"])
        if forme.get("dash"):
            ln = objet.line._get_or_add_ln()
            for ancien in ln.findall(qn("a:prstDash")):
                ln.remove(ancien)
            tiret = ln.makeelement(qn("a:prstDash"), {"val": "dash"})
            ln.append(tiret)
    else:
        objet.line.fill.background()

    sans_ombre(objet)
    cadre = objet.text_frame
    if forme["texte"]:
        poser_texte(cadre, forme, police)
    else:
        cadre.text = ""
    if forme.get("lien") is not None and forme["lien"] in cibles:
        objet.click_action.target_slide = cibles[forme["lien"]]
    return objet


def ajouter_texte(diapo, forme, police):
    objet = diapo.shapes.add_textbox(emu(forme["x"]), emu(forme["y"]),
                                     emu(forme["w"]), emu(forme["h"]))
    poser_texte(objet.text_frame, forme, police)
    sans_ombre(objet)
    return objet


GABARIT_BEZIER = (
    '<p:sp %s>'
    '<p:nvSpPr><p:cNvPr id="%%(id)d" name="%%(nom)s"/>'
    '<p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
    '<p:spPr>'
    '<a:xfrm><a:off x="%%(x)d" y="%%(y)d"/><a:ext cx="%%(cx)d" cy="%%(cy)d"/></a:xfrm>'
    '<a:custGeom><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/>'
    '<a:rect l="0" t="0" r="%%(cx)d" b="%%(cy)d"/>'
    '<a:pathLst><a:path w="%%(cx)d" h="%%(cy)d">'
    '<a:moveTo><a:pt x="%%(x0)d" y="%%(y0)d"/></a:moveTo>'
    '<a:cubicBezTo><a:pt x="%%(x1)d" y="%%(y1)d"/>'
    '<a:pt x="%%(x2)d" y="%%(y2)d"/><a:pt x="%%(x3)d" y="%%(y3)d"/></a:cubicBezTo>'
    '</a:path></a:pathLst></a:custGeom>'
    '<a:noFill/>'
    '<a:ln w="%%(lw)d" cap="rnd"><a:solidFill><a:srgbClr val="%%(couleur)s"/></a:solidFill>'
    '%%(dash)s%%(fleche)s<a:round/></a:ln>'
    '<a:effectLst/>'
    '</p:spPr>'
    '<p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody>'
    '</p:sp>') % nsdecls("p", "a")


def ajouter_bezier(diapo, forme, identifiant):
    """Une seule courbe de Bézier cubique par flux, en géométrie personnalisée.

    Les points d'un `custGeom` sont relatifs à la boîte englobante de la
    courbe, pas à la planche : d'où le décalage appliqué ci-dessous.
    """
    points = forme["pts"]
    marge = forme["lw"] * 3.0
    xs = [p[0] for p in points]
    ys = [p[1] for p in points]
    x0, y0 = min(xs) - marge, min(ys) - marge
    largeur = max(2.0, max(xs) - min(xs) + 2 * marge)
    hauteur = max(2.0, max(ys) - min(ys) + 2 * marge)
    relatifs = []
    for px_, py_ in points:
        relatifs.append((int(round((px_ - x0) * EMU_PAR_POINT)),
                         int(round((py_ - y0) * EMU_PAR_POINT))))
    valeurs = {
        "id": identifiant, "nom": "Flux %d" % identifiant,
        "x": int(round(x0 * EMU_PAR_POINT)), "y": int(round(y0 * EMU_PAR_POINT)),
        "cx": int(round(largeur * EMU_PAR_POINT)),
        "cy": int(round(hauteur * EMU_PAR_POINT)),
        "lw": int(round(forme["lw"] * EMU_PAR_POINT)),
        "couleur": forme["couleur"].lstrip("#").upper(),
        "dash": '<a:prstDash val="dash"/>' if forme.get("dash") else "",
        "fleche": ('<a:tailEnd type="triangle" w="med" len="med"/>'
                   if forme.get("fleche") else ""),
    }
    for rang, (rx, ry) in enumerate(relatifs):
        valeurs["x%d" % rang] = rx
        valeurs["y%d" % rang] = ry
    from lxml import etree
    element = etree.fromstring(GABARIT_BEZIER % valeurs)
    diapo.shapes._spTree.append(element)


# ═════════════════════════════════════════════════════════════════════════════
# Graphiques natifs
# ═════════════════════════════════════════════════════════════════════════════

# Ordre imposé par le schéma DrawingML pour les enfants de c:catAx
FIN_AXE = ("c:tickLblSkip", "c:tickMarkSkip", "c:noMultiLvlLbl", "c:extLst")


def forcer_toutes_les_categories(axe):
    """Affiche une étiquette par catégorie.

    Sans cela, le rendu n'en garde qu'une sur deux dès que les libellés sont
    serrés, et la moitié des services disparaît du classement.
    """
    element = axe._element
    for balise, valeur in (("c:tickLblSkip", "1"), ("c:tickMarkSkip", "1")):
        for ancien in element.findall(qn(balise)):
            element.remove(ancien)
        noeud = element.makeelement(qn(balise), {"val": valeur})
        suivant = None
        for candidat in FIN_AXE[FIN_AXE.index(balise) + 1:]:
            trouve = element.find(qn(candidat))
            if trouve is not None:
                suivant = trouve
                break
        if suivant is None:
            element.append(noeud)
        else:
            suivant.addprevious(noeud)


def ajouter_graphe(diapo, forme, vue, charte, cibles, liens):
    """Objet graphique natif : les données restent éditables dans PowerPoint."""
    spec = forme["spec"]
    donnees = CategoryChartData()
    if spec["genre"] == "doughnut":
        valeurs = list(vue[spec["donnees"]].values())
        donnees.categories = spec["libelles"]
        donnees.add_series("Effectif", valeurs)
        genre = XL_CHART_TYPE.DOUGHNUT
    else:
        # PowerPoint trace la première catégorie en bas : on inverse l'ordre
        # pour retrouver le classement décroissant de haut en bas du HTML
        classement = list(reversed(vue["classement"]))
        donnees.categories = [ligne["court"] for ligne in classement]
        donnees.add_series(spec["titre"],
                           [ligne[spec["champ"]] for ligne in classement])
        genre = XL_CHART_TYPE.BAR_CLUSTERED

    cadre = diapo.shapes.add_chart(genre, emu(forme["x"]), emu(forme["y"]),
                                   emu(forme["w"]), emu(forme["h"]), donnees)
    graphique = cadre.chart
    # le titre du graphique doublerait celui déjà posé par layout.py
    graphique.has_title = False
    graphique.font.size = Pt(6.5)
    graphique.font.name = charte["police_pptx"]
    graphique.font.color.rgb = rgb(charte["gris"])

    if spec["genre"] == "doughnut":
        graphique.has_legend = True
        graphique.legend.position = XL_LEGEND_POSITION.RIGHT
        graphique.legend.include_in_layout = False
        points = graphique.plots[0].series[0].points
        for rang, couleur in enumerate(spec["couleurs"]):
            if rang < len(points):
                points[rang].format.fill.solid()
                points[rang].format.fill.fore_color.rgb = rgb(couleur)
                points[rang].format.line.color.rgb = rgb(charte["blanc"])
                points[rang].format.line.width = emu(1.4)
    else:
        graphique.has_legend = False
        serie = graphique.plots[0].series[0]
        serie.format.fill.solid()
        serie.format.fill.fore_color.rgb = rgb(spec["couleurs"][0])
        serie.format.line.fill.background()
        graphique.category_axis.has_major_gridlines = False
        graphique.value_axis.has_major_gridlines = True
        forcer_toutes_les_categories(graphique.category_axis)
    return cadre


# ═════════════════════════════════════════════════════════════════════════════
# Assemblage
# ═════════════════════════════════════════════════════════════════════════════

def preparer(presentation):
    """Format 16:9 natif et gabarit vierge."""
    presentation.slide_width = emu(layout.PAGE_W)
    presentation.slide_height = emu(layout.PAGE_H)
    return presentation.slide_layouts[6]      # « Vide »


def generer(carto, chemin):
    pages = layout.construire(carto)
    charte = carto["charte"]
    police = charte["police_pptx"]
    vue = model.indicateurs(carto, model.TOUS)
    liens = {s["id"]: rang + 1 for rang, s in enumerate(carto["services"])}

    presentation = Presentation()
    vierge = preparer(presentation)
    diapos = [presentation.slides.add_slide(vierge) for _ in pages]
    cibles = {indice: diapo for indice, diapo in enumerate(diapos)}

    identifiant = 5000
    for page, diapo in zip(pages, diapos):
        # fond blanc explicite : le thème pose sinon un fond de gabarit
        fond = diapo.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0,
                                      emu(layout.PAGE_W), emu(layout.PAGE_H))
        fond.fill.solid()
        fond.fill.fore_color.rgb = rgb(charte["blanc"])
        fond.line.fill.background()
        sans_ombre(fond)
        fond.text_frame.text = ""

        for forme in page["formes"]:
            if forme["type"] == "rect":
                ajouter_rect(diapo, forme, police, cibles)
            elif forme["type"] == "texte":
                ajouter_texte(diapo, forme, police)
            elif forme["type"] == "bezier":
                identifiant += 1
                ajouter_bezier(diapo, forme, identifiant)
            elif forme["type"] == "graphe":
                ajouter_graphe(diapo, forme, vue, charte, cibles, liens)

        diapo.notes_slide.notes_text_frame.text = (
            "%s — page %d sur %d. Généré depuis data/carto.json ; toute "
            "modification de contenu passe par ce fichier."
            % (page["titre"], page["indice"] + 1, len(pages)))

    chemin.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(chemin))
    return pages, diapos


def main():
    analyseur = argparse.ArgumentParser(description=__doc__)
    analyseur.add_argument("--out", default=str(SORTIE),
                           help="fichier PPTX à écrire")
    arguments = analyseur.parse_args()

    carto = model.charger()
    chemin = Path(arguments.out)
    pages, diapos = generer(carto, chemin)
    formes = sum(len(p["formes"]) for p in pages)
    print("%s : %d diapositives, %d formes natives, %.0f Ko"
          % (chemin, len(pages), formes, chemin.stat().st_size / 1024.0))


if __name__ == "__main__":
    main()
